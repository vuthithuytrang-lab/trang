"""Lọc các bài cần tạo HTML từ file inlink (bản xuất .xlsx).

python3 loc_bai.py <file.xlsx> <viec_can_lam.json>

Điều kiện: cột "Tình trạng" khác "Đã hoàn thành", cột "Chị Lộc duyệt" có chữ "ok",
cột "HTML" đang trống. Cột được dò theo tên tiêu đề nên chuyển chỗ cột vẫn chạy đúng.
"""
import sys, json, re
import openpyxl

SHEET = 'Plan internal links - Tổng'

def col_letter(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb[SHEET]
head = {str(ws.cell(1, c).value or '').strip(): c for c in range(1, ws.max_column + 1)}

def find(*names):
    for n in names:
        if n in head: return head[n]
    raise SystemExit(f'Không thấy cột tiêu đề: {names}')

cB, cD, cK, cL = find('URL'), find('Tình trạng'), find('FILE BLOG'), find('Chị Lộc duyệt')
cI, cJ = find('Anchor text (Trang tỷ giá)'), find('URL đích (Trang tỷ giá)')
cH = find('HTML')
pairs_cols = []
for n in range(1, 6):
    a = next((head[k] for k in head if re.fullmatch(rf'Anchor ?text {n}', k, re.I)), None)
    u = head.get(f'URL đích {n}')
    if a and u: pairs_cols.append((a, u))

g = lambda r, c: ws.cell(r, c).value
jobs, skipped = [], []
for r in range(2, ws.max_row + 1):
    url = str(g(r, cB) or '').strip()
    d = str(g(r, cD) or '').strip()
    l = str(g(r, cL) or '').strip()
    if not url or d == 'Đã hoàn thành' or 'ok' not in l.lower() or str(g(r, cH) or '').strip():
        continue
    m = re.search(r'/d/([\w-]+)', str(g(r, cK) or ''))
    if not m:
        skipped.append(r); continue
    jobs.append(dict(row=r, url=url, docId=m.group(1), I=g(r, cI), J=g(r, cJ),
                     note='' if l.lower() == 'ok' else l,
                     pairs=[(g(r, a), g(r, u)) for a, u in pairs_cols if g(r, a) and g(r, u)],
                     html_col=col_letter(cH), sheet=SHEET))
json.dump(jobs, open(sys.argv[2], 'w', encoding='utf-8'), ensure_ascii=False, default=str)
print('Cột HTML:', col_letter(cH), '| Số bài cần chạy:', len(jobs), [j['row'] for j in jobs])
if skipped:
    print('Bỏ qua (cột FILE BLOG không có link Docs):', skipped)
